"""
scripts/reporting/export_test_report.py

Generates an Excel test report (Test Case Checklist) from the pytest suite.

Usage (from repo root, venv activated):
    .venv\\Scripts\\python.exe scripts\\reporting\\export_test_report.py

Requires: openpyxl
    pip install openpyxl

Output: docs/test_report.xlsx
    - Summary sheet: pass/fail totals
    - Automated sheet: one row per pytest test case (auto-filled from
      junit-xml results + docstrings), color-coded by status
    - Manual sheet: empty template with the same columns, for manually
      executed test cases (e.g. SSH remote-collection live tests, UI
      click-through tests) that aren't covered by pytest
"""
import ast
import subprocess
import sys
import xml.etree.ElementTree as ET
from pathlib import Path
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]  # repo root (scripts/reporting/../..)
JUNIT_PATH = ROOT / "scripts" / "reporting" / "_junit_results.xml"

HEADERS = [
    "Test Case ID", "Module", "Test Name", "Objective",
    "Preconditions", "Expected Result", "Actual Result", "Status", "Duration (s)",
]
COLUMN_WIDTHS = [12, 40, 45, 45, 20, 30, 40, 10, 12]


def run_pytest_junit() -> None:
    """Run the full suite and produce a junit-xml report."""
    JUNIT_PATH.parent.mkdir(parents=True, exist_ok=True)
    print("Running full pytest suite (this may take a minute)...")
    subprocess.run(
        [sys.executable, "-m", "pytest", "-q", f"--junitxml={JUNIT_PATH}"],
        cwd=ROOT,
        check=False,  # don't crash on the known intentional TDD failure
    )


def extract_docstrings(test_file: Path) -> dict:
    """Map test function name -> first line of its docstring (if any)."""
    try:
        tree = ast.parse(test_file.read_text(encoding="utf-8"))
    except (SyntaxError, UnicodeDecodeError, OSError):
        return {}
    result = {}
    for node in ast.walk(tree):
        if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)) and node.name.startswith("test_"):
            doc = ast.get_docstring(node)
            result[node.name] = doc.strip().splitlines()[0] if doc else ""
    return result


def humanize(name: str) -> str:
    return name.replace("test_", "", 1).replace("_", " ").strip().capitalize()


def parse_junit(junit_path: Path) -> list:
    tree = ET.parse(junit_path)
    root = tree.getroot()
    rows = []
    for testcase in root.iter("testcase"):
        classname = testcase.get("classname", "")
        name = testcase.get("name", "")
        time = float(testcase.get("time", 0))
        status = "Pass"
        message = ""
        failure = testcase.find("failure")
        error = testcase.find("error")
        skipped = testcase.find("skipped")
        if failure is not None:
            status = "Fail"
            message = (failure.get("message") or "")[:300]
        elif error is not None:
            status = "Error"
            message = (error.get("message") or "")[:300]
        elif skipped is not None:
            status = "Skipped"
            message = (skipped.get("message") or "")[:300]
        rows.append({
            "module": classname,
            "test_name": name,
            "duration_s": round(time, 3),
            "status": status,
            "message": message,
        })
    return rows


def build_workbook(rows: list, docstrings_by_module: dict) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Automated"

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    pass_fill = PatternFill("solid", fgColor="C6EFCE")
    fail_fill = PatternFill("solid", fgColor="FFC7CE")
    skip_fill = PatternFill("solid", fgColor="FFEB9C")

    def write_header(sheet):
        sheet.append(HEADERS)
        for col_idx in range(1, len(HEADERS) + 1):
            cell = sheet.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for idx, w in enumerate(COLUMN_WIDTHS, start=1):
            sheet.column_dimensions[get_column_letter(idx)].width = w
        sheet.freeze_panes = "A2"

    write_header(ws)

    for i, row in enumerate(rows, start=1):
        module = row["module"]
        docmap = docstrings_by_module.get(module, {})
        objective = docmap.get(row["test_name"]) or humanize(row["test_name"])
        expected = "Assertions pass without error"
        actual = row["message"] if row["status"] in ("Fail", "Error") else "As expected"

        ws.append([
            f"TC-{i:04d}",
            module,
            row["test_name"],
            objective,
            "",  # Preconditions - fill in manually where relevant
            expected,
            actual,
            row["status"],
            row["duration_s"],
        ])

        status_cell = ws.cell(row=i + 1, column=8)
        if row["status"] == "Pass":
            status_cell.fill = pass_fill
        elif row["status"] in ("Fail", "Error"):
            status_cell.fill = fail_fill
        else:
            status_cell.fill = skip_fill

    # Manual test case template sheet
    ws2 = wb.create_sheet("Manual")
    write_header(ws2)

    # Summary sheet (placed first)
    ws3 = wb.create_sheet("Summary", 0)
    total = len(rows)
    passed = sum(1 for r in rows if r["status"] == "Pass")
    failed = sum(1 for r in rows if r["status"] in ("Fail", "Error"))
    skipped = sum(1 for r in rows if r["status"] == "Skipped")
    ws3.append(["CABTA Automated Test Report"])
    ws3.append(["Generated", datetime.now().strftime("%Y-%m-%d %H:%M")])
    ws3.append([])
    ws3.append(["Total", total])
    ws3.append(["Passed", passed])
    ws3.append(["Failed / Error", failed])
    ws3.append(["Skipped", skipped])
    ws3["A1"].font = Font(bold=True, size=14)
    ws3.column_dimensions["A"].width = 20
    ws3.column_dimensions["B"].width = 30

    return wb


def main() -> None:
    run_pytest_junit()
    if not JUNIT_PATH.exists():
        print("junit xml was not produced — check the pytest run above for errors.")
        return

    rows = parse_junit(JUNIT_PATH)

    docstrings_by_module = {}
    for row in rows:
        module = row["module"]
        if module in docstrings_by_module:
            continue
        # pytest classname is dotted, e.g. "tests.test_chat_structured_params"
        test_file = ROOT / (module.replace(".", "/") + ".py")
        if test_file.exists():
            docstrings_by_module[module] = extract_docstrings(test_file)

    wb = build_workbook(rows, docstrings_by_module)
    out_path = ROOT / "docs" / "test_report.xlsx"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)

    total = len(rows)
    passed = sum(1 for r in rows if r["status"] == "Pass")
    print(f"Saved: {out_path}")
    print(f"{passed}/{total} passed")


if __name__ == "__main__":
    main()